"""Ограниченные ZIP/GZIP и Extract from File runtime-адаптеры."""

import base64
import csv
import gzip
import html
import io
import json
import zipfile
from collections.abc import Mapping, Sequence
from datetime import date, datetime, time
from html.parser import HTMLParser
from typing import Literal

import openpyxl
import xlrd
import xlwt
from pydantic import Field, field_validator
from pypdf import PdfReader

from ..domain.models.workflow import WorkflowNode
from ..runtime.executors import NodeExecutionContext
from ..runtime.models import NodeResult, WorkflowItem
from .errors import UnsupportedNodeConfigurationError
from .expressions import N8nExpressionEvaluator
from .nodes import EmptyOptions, _StrictParameters

_MAX_FILE_BYTES = 50 * 1024 * 1024
_MAX_CSV_ROWS = 50_000
_MAX_SPREADSHEET_CELLS = 2_000_000
_XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XLS_MIME_TYPE = "application/vnd.ms-excel"


class CompressionParameters(_StrictParameters):
    """Поддерживаемые настройки ZIP/GZIP узла Compression."""

    operation: Literal["compress", "decompress"]
    binary_property_name: str = Field(default="data", alias="binaryPropertyName")
    output_format: Literal["zip", "gzip"] = Field(default="zip", alias="outputFormat")
    file_name: str = Field(default="data", alias="fileName", min_length=1)
    binary_property_output: str = Field(default="data", alias="binaryPropertyOutput")
    output_prefix: str = Field(default="file", alias="outputPrefix")
    options: EmptyOptions = Field(default_factory=EmptyOptions)


class ExtractOptions(_StrictParameters):
    """Параметры чтения таблиц, текста и PDF."""

    encoding: Literal["utf-8", "utf-8-sig", "latin-1", "ascii"] = "utf-8-sig"
    strip_bom: bool = Field(default=True, alias="stripBOM")
    keep_source: bool | Literal["none", "json", "binary", "both"] = Field(
        default="json", alias="keepSource"
    )
    header_row: bool = Field(default=True, alias="headerRow")
    include_empty_cells: bool = Field(default=True, alias="includeEmptyCells")
    delimiter: str = Field(default=",", min_length=1, max_length=1)
    sheet_name: str | None = Field(default=None, alias="sheetName", min_length=1)
    join_pages: bool = Field(default=True, alias="joinPages")
    max_pages: int = Field(default=500, alias="maxPages", gt=0, le=500, strict=True)

    @field_validator("encoding", mode="before")
    @classmethod
    def normalize_encoding(cls, value: object) -> object:
        """Принимать краткие имена кодировок из n8n export."""
        if isinstance(value, str):
            return {"utf8": "utf-8", "latin1": "latin-1"}.get(value.lower(), value.lower())
        return value


class ExtractFileParameters(_StrictParameters):
    """Настройки извлечения таблиц и текста из binary-поля item."""

    operation: Literal["csv", "html", "fromJson", "pdf", "text", "xlsx", "xls", "binaryToPropery"]
    binary_property_name: str = Field(default="data", alias="binaryPropertyName")
    destination_key: str = Field(default="row", alias="destinationKey")
    options: ExtractOptions = Field(default_factory=ExtractOptions)


class ConvertToFileOptions(_StrictParameters):
    """Параметры генерации CSV, HTML, TXT и Excel-файлов."""

    file_name: str = Field(default="", alias="fileName")
    encoding: Literal["utf-8", "latin-1", "ascii"] = "utf-8"
    add_bom: bool = Field(default=False, alias="addBOM")
    header_row: bool = Field(default=True, alias="headerRow")
    delimiter: str = Field(default=",", min_length=1, max_length=1)
    compression: bool = False
    sheet_name: str = Field(default="Sheet1", alias="sheetName", min_length=1)

    @field_validator("encoding", mode="before")
    @classmethod
    def normalize_encoding(cls, value: object) -> object:
        """Принимать краткие имена кодировок из n8n export."""
        if isinstance(value, str):
            return {"utf8": "utf-8", "latin1": "latin-1"}.get(value.lower(), value.lower())
        return value


class ConvertToFileParameters(_StrictParameters):
    """Параметры преобразования JSON items в файл для binary-поля."""

    operation: Literal["csv", "html", "toText", "xlsx", "xls"]
    source_property: str = Field(default="data", alias="sourceProperty")
    binary_property_name: str = Field(default="data", alias="binaryPropertyName")
    options: ConvertToFileOptions = Field(default_factory=ConvertToFileOptions)


def _binary_bytes(item: WorkflowItem, property_name: str) -> tuple[bytes, dict[str, object]]:
    """Прочитать n8n-совместимое base64 binary-поле и проверить размер."""
    binary = (item.binary or {}).get(property_name)
    if not isinstance(binary, dict) or not isinstance(binary.get("data"), str):
        raise UnsupportedNodeConfigurationError(
            f"Входной item не содержит binary-поле {property_name!r} с base64-данными."
        )
    try:
        content = base64.b64decode(binary["data"], validate=True)
    except (ValueError, base64.binascii.Error) as error:
        raise UnsupportedNodeConfigurationError(
            f"Binary-поле {property_name!r} содержит некорректный base64."
        ) from error
    if len(content) > _MAX_FILE_BYTES:
        raise UnsupportedNodeConfigurationError("Размер входного файла превышает 50 MiB.")
    return content, binary


def _binary_metadata(content: bytes, name: str, mime_type: str) -> dict[str, object]:
    """Сформировать inline binary-значение для сериализуемого WorkflowItem."""
    if len(content) > _MAX_FILE_BYTES:
        raise UnsupportedNodeConfigurationError("Созданный файл превышает лимит 50 MiB.")
    return {
        "data": base64.b64encode(content).decode("ascii"),
        "fileName": name,
        "fileExtension": name.rsplit(".", 1)[-1] if "." in name else "",
        "mimeType": mime_type,
    }


class _FirstHtmlTableParser(HTMLParser):
    """Собрать ячейки первой HTML-таблицы, не исполняя встроенный JavaScript."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self._table_depth = 0
        self._found_table = False
        self._row: list[str] | None = None
        self._cell_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag == "table":
            if not self._found_table:
                self._found_table = True
                self._table_depth = 1
            elif self._table_depth:
                self._table_depth += 1
        elif self._table_depth == 1 and tag == "tr":
            self._row = []
        elif self._table_depth == 1 and self._row is not None and tag in {"th", "td"}:
            self._cell_parts = []

    def handle_endtag(self, tag: str) -> None:
        if self._table_depth == 1 and tag in {"th", "td"} and self._cell_parts is not None:
            self._row = self._row or []
            self._row.append("".join(self._cell_parts).strip())
            self._cell_parts = None
        elif self._table_depth == 1 and tag == "tr" and self._row is not None:
            if self._row:
                if len(self.rows) >= _MAX_CSV_ROWS + 1:
                    raise UnsupportedNodeConfigurationError(
                        "HTML-таблица содержит больше 50 000 строк."
                    )
                self.rows.append(self._row)
            self._row = None
        elif tag == "table" and self._table_depth:
            self._table_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._cell_parts is not None:
            self._cell_parts.append(data)


def _rows_to_records(
    rows: Sequence[Sequence[object]], *, header_row: bool, include_empty_cells: bool
) -> list[dict[str, object]]:
    """Преобразовать строки таблицы в JSON-совместимые записи."""
    if not rows:
        return []
    width = max(len(row) for row in rows)
    if width * len(rows) > _MAX_SPREADSHEET_CELLS:
        raise UnsupportedNodeConfigurationError("Таблица превышает лимит в 2 000 000 ячеек.")

    body = rows[1:] if header_row else rows
    if header_row:
        raw_headers = list(rows[0])
        headers: list[str] = []
        seen: dict[str, int] = {}
        for column_index in range(width):
            candidate = (
                str(raw_headers[column_index]).strip() if column_index < len(raw_headers) else ""
            )
            base_name = candidate or str(column_index)
            seen[base_name] = seen.get(base_name, 0) + 1
            headers.append(base_name if seen[base_name] == 1 else f"{base_name}_{seen[base_name]}")
    else:
        headers = [str(column_index) for column_index in range(width)]

    records: list[dict[str, object]] = []
    for row in body:
        if len(records) >= _MAX_CSV_ROWS:
            raise UnsupportedNodeConfigurationError("Таблица содержит больше 50 000 строк.")
        record: dict[str, object] = {}
        for column_index, key in enumerate(headers):
            value = row[column_index] if column_index < len(row) else None
            if isinstance(value, (datetime, date, time)):
                value = value.isoformat()
            if not include_empty_cells and value in (None, ""):
                continue
            record[key] = value
        records.append(record)
    return records


def _text_encoding(encoding: str) -> str:
    """Map common n8n encoding names to codecs understood by Python."""
    return {"utf8": "utf-8", "latin1": "latin-1"}.get(encoding.lower(), encoding.lower())


def _decode_file_text(content: bytes, encoding: str, strip_bom: bool) -> str:
    """Decode a text-like file and optionally remove its Unicode BOM."""
    codec = _text_encoding(encoding)
    if codec == "utf-8-sig" and not strip_bom:
        codec = "utf-8"
    try:
        text = content.decode(codec)
    except (LookupError, UnicodeDecodeError) as error:
        raise UnsupportedNodeConfigurationError(
            "Не удалось прочитать файл в выбранной кодировке."
        ) from error
    return text.removeprefix("\ufeff") if strip_bom else text


def _lookup_json_path(value: object, path: str) -> object:
    """Resolve a dot-separated field path in an input item's JSON object."""
    current = value
    for part in path.split("."):
        if not isinstance(current, Mapping) or part not in current:
            raise UnsupportedNodeConfigurationError(
                f"В item отсутствует поле {path!r}, указанное в sourceProperty."
            )
        current = current[part]
    return current


def _set_json_path(target: dict[str, object], path: str, value: object) -> None:
    """Assign an output to a dot-separated nested JSON key."""
    parts = path.split(".")
    current = target
    for part in parts[:-1]:
        child = current.get(part)
        if not isinstance(child, dict):
            child = {}
            current[part] = child
        current = child
    current[parts[-1]] = value


def _spreadsheet_rows(items: Sequence[WorkflowItem]) -> tuple[list[str], list[list[object]]]:
    """Create stable table columns and rows from n8n JSON items."""
    if len(items) > _MAX_CSV_ROWS:
        raise UnsupportedNodeConfigurationError("Таблица содержит больше 50 000 строк.")
    records: list[Mapping[str, object]] = []
    columns: list[str] = []
    for item in items:
        if not isinstance(item.json_data, Mapping):
            raise UnsupportedNodeConfigurationError(
                "Convert to File для табличного формата ожидает JSON-object на входе."
            )
        record = dict(item.json_data)
        records.append(record)
        for key in record:
            if str(key) not in columns:
                columns.append(str(key))
    if not columns:
        raise UnsupportedNodeConfigurationError("Нельзя создать таблицу без полей JSON-object.")
    rows = [[record.get(key) for key in columns] for record in records]
    return columns, rows


def _spreadsheet_value(value: object) -> object:
    """Convert nested values to text while preserving simple Excel cell types."""
    if isinstance(value, (str, int, float, bool, date, datetime, time)) or value is None:
        return value
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _file_name(options_name: str, extension: str) -> str:
    """Use the configured output name or a predictable format-specific default."""
    return options_name or f"data.{extension}"


def _with_extension(name: str, extension: str) -> str:
    """Add a default suffix only when the configured name has no extension."""
    return name if "." in name.rsplit("/", 1)[-1] else f"{name}.{extension}"


class CompressionExecutor:
    """Сжимать или распаковывать один binary-файл без записи на диск."""

    is_trigger = False

    def __init__(self, parameters: CompressionParameters):
        self._parameters = parameters
        self._expressions = N8nExpressionEvaluator()

    async def execute(
        self,
        node: WorkflowNode,
        items: tuple[WorkflowItem, ...],
        context: NodeExecutionContext,
    ) -> NodeResult:
        output: list[WorkflowItem] = []
        paired_input_indexes: list[int] = []
        for item_index, item in enumerate(items):
            content, input_metadata = _binary_bytes(item, self._parameters.binary_property_name)
            if self._parameters.operation == "compress":
                archive_name = self._expressions.evaluate(
                    self._parameters.file_name,
                    item,
                    item_index=item_index,
                    context=context,
                )
                if not isinstance(archive_name, str) or not archive_name:
                    raise UnsupportedNodeConfigurationError(
                        "Compression fileName должен вычисляться в непустую строку."
                    )
                input_file_name = input_metadata.get("fileName")
                if not isinstance(input_file_name, str) or not input_file_name:
                    input_file_name = self._parameters.binary_property_name
                archive = self._compress(content, input_file_name)
                name = self._compressed_name(archive_name)
                mimetype = (
                    "application/zip"
                    if self._parameters.output_format == "zip"
                    else "application/gzip"
                )
                output.append(
                    _replace_binary(
                        item,
                        self._parameters.binary_property_output,
                        _binary_metadata(archive, name, mimetype),
                    )
                )
                paired_input_indexes.append(item_index)
            else:
                for extracted_name, extracted, mime_type in self._decompress(content):
                    output_property = (
                        self._parameters.output_prefix
                        if self._parameters.output_format == "gzip"
                        else f"{self._parameters.output_prefix}{extracted_name}"
                    )
                    output.append(
                        _replace_binary(
                            item,
                            output_property,
                            _binary_metadata(extracted, extracted_name, mime_type),
                        )
                    )
                    paired_input_indexes.append(item_index)
        return NodeResult(
            outputs=(tuple(output),),
            paired_input_indexes=(tuple(paired_input_indexes),),
        )

    def _compressed_name(self, file_name: str) -> str:
        name = file_name
        if self._parameters.output_format == "gzip" and not name.endswith(".gz"):
            return f"{name}.gz"
        if self._parameters.output_format == "zip" and not name.endswith(".zip"):
            return f"{name}.zip"
        return name

    def _compress(self, content: bytes, input_file_name: str) -> bytes:
        if self._parameters.output_format == "gzip":
            return gzip.compress(content)
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(input_file_name, content)
        return buffer.getvalue()

    def _decompress(self, content: bytes) -> list[tuple[str, bytes, str]]:
        if self._parameters.output_format == "gzip":
            try:
                with gzip.GzipFile(fileobj=io.BytesIO(content)) as stream:
                    extracted = stream.read(_MAX_FILE_BYTES + 1)
            except (OSError, EOFError) as error:
                raise UnsupportedNodeConfigurationError(
                    "Не удалось распаковать GZIP-файл."
                ) from error
            if len(extracted) > _MAX_FILE_BYTES:
                raise UnsupportedNodeConfigurationError("Распакованный файл превышает 50 MiB.")
            return [(self._gzip_output_name(), extracted, "application/octet-stream")]

        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                files = [entry for entry in archive.infolist() if not entry.is_dir()]
                if not files:
                    raise UnsupportedNodeConfigurationError("ZIP-архив не содержит файлов.")
                extracted_files: list[tuple[str, bytes, str]] = []
                total_size = 0
                for entry in files:
                    total_size += entry.file_size
                    if total_size > _MAX_FILE_BYTES:
                        raise UnsupportedNodeConfigurationError(
                            "Общий размер распакованных файлов превышает 50 MiB."
                        )
                    with archive.open(entry) as stream:
                        data = stream.read(_MAX_FILE_BYTES + 1)
                    if len(data) != entry.file_size or len(data) > _MAX_FILE_BYTES:
                        raise UnsupportedNodeConfigurationError(
                            "ZIP-архив содержит некорректный или слишком большой файл."
                        )
                    extracted_files.append(
                        (entry.filename.rsplit("/", 1)[-1], data, "application/octet-stream")
                    )
                return extracted_files
        except (OSError, zipfile.BadZipFile, RuntimeError) as error:
            raise UnsupportedNodeConfigurationError("Не удалось распаковать ZIP-файл.") from error

    def _gzip_output_name(self) -> str:
        return self._parameters.file_name.removesuffix(".gz")


class ConvertToFileExecutor:
    """Создать CSV, HTML, TXT, XLS или XLSX в inline binary-поле item."""

    is_trigger = False

    _FORMATS = {
        "csv": ("csv", "text/csv"),
        "html": ("html", "text/html"),
        "toText": ("txt", "text/plain"),
        "xls": ("xls", _XLS_MIME_TYPE),
        "xlsx": ("xlsx", _XLSX_MIME_TYPE),
    }

    def __init__(self, parameters: ConvertToFileParameters):
        self._parameters = parameters

    async def execute(
        self,
        node: WorkflowNode,
        items: tuple[WorkflowItem, ...],
        context: NodeExecutionContext,
    ) -> NodeResult:
        del node, context
        if not items:
            return NodeResult(outputs=((),), paired_input_indexes=((),))
        if self._parameters.operation == "toText":
            output = [self._convert_text(item) for item in items]
            return NodeResult(
                outputs=(tuple(output),),
                paired_input_indexes=(tuple(range(len(items))),),
            )

        content, file_name, mime_type = self._convert_table(items)
        output = WorkflowItem(
            json={},
            binary={
                self._parameters.binary_property_name: _binary_metadata(
                    content, file_name, mime_type
                )
            },
        )
        paired_input = 0 if len(items) == 1 else None
        return NodeResult(
            outputs=((output,),),
            paired_input_indexes=((paired_input,),),
        )

    def _convert_text(self, item: WorkflowItem) -> WorkflowItem:
        value = _lookup_json_path(item.json_data, self._parameters.source_property)
        if not isinstance(value, str):
            raise UnsupportedNodeConfigurationError(
                "Convert to File с operation=toText ожидает строку в sourceProperty."
            )
        try:
            content = value.encode(_text_encoding(self._parameters.options.encoding))
        except (LookupError, UnicodeEncodeError) as error:
            raise UnsupportedNodeConfigurationError(
                "Текст не помещается в выбранную кодировку."
            ) from error
        if self._parameters.options.add_bom and self._parameters.options.encoding == "utf-8":
            content = b"\xef\xbb\xbf" + content
        file_name = _with_extension(_file_name(self._parameters.options.file_name, "txt"), "txt")
        return WorkflowItem(
            json={},
            binary={
                self._parameters.binary_property_name: _binary_metadata(
                    content, file_name, "text/plain"
                )
            },
        )

    def _convert_table(self, items: tuple[WorkflowItem, ...]) -> tuple[bytes, str, str]:
        operation = self._parameters.operation
        extension, mime_type = self._FORMATS[operation]
        file_name = _with_extension(
            _file_name(self._parameters.options.file_name, extension), extension
        )
        columns, rows = _spreadsheet_rows(items)
        if self._parameters.options.header_row:
            rows.insert(0, columns)
        rows = [[_spreadsheet_value(value) for value in row] for row in rows]

        if operation == "csv":
            text_buffer = io.StringIO(newline="")
            csv.writer(
                text_buffer,
                delimiter=self._parameters.options.delimiter,
                lineterminator="\n",
            ).writerows(rows)
            try:
                content = text_buffer.getvalue().encode(
                    _text_encoding(self._parameters.options.encoding)
                )
            except (LookupError, UnicodeEncodeError) as error:
                raise UnsupportedNodeConfigurationError(
                    "Таблица не помещается в выбранную кодировку."
                ) from error
            if self._parameters.options.add_bom and self._parameters.options.encoding == "utf-8":
                content = b"\xef\xbb\xbf" + content
            return content, file_name, mime_type

        if operation == "html":
            content = self._render_html(rows).encode("utf-8")
            return content, file_name, mime_type

        if self._parameters.options.compression and operation != "xlsx":
            raise UnsupportedNodeConfigurationError(
                "Опция compression поддерживается только при создании XLSX."
            )
        content = self._write_excel(rows, operation)
        return content, file_name, mime_type

    def _render_html(self, rows: Sequence[Sequence[object]]) -> str:
        """Render a safe UTF-8 HTML table from JSON rows."""
        table_rows = []
        for row_index, row in enumerate(rows):
            tag = "th" if self._parameters.options.header_row and row_index == 0 else "td"
            cells = "".join(
                f"<{tag}>{html.escape('' if value is None else str(value))}</{tag}>"
                for value in row
            )
            table_rows.append(f"<tr>{cells}</tr>")
        return (
            '<!doctype html><html><head><meta charset="utf-8"></head><body><table>'
            + "".join(table_rows)
            + "</table></body></html>"
        )

    def _write_excel(self, rows: Sequence[Sequence[object]], operation: str) -> bytes:
        sheet_name = self._parameters.options.sheet_name
        if any(character in sheet_name for character in "[]:*?/\\") or len(sheet_name) > 31:
            raise UnsupportedNodeConfigurationError("Некорректное имя листа Excel.")
        buffer = io.BytesIO()
        try:
            if operation == "xlsx":
                workbook = openpyxl.Workbook(write_only=True)
                sheet = workbook.create_sheet(title=sheet_name)
                for row in rows:
                    sheet.append(list(row))
                workbook.save(buffer)
                workbook.close()
                compression = (
                    zipfile.ZIP_DEFLATED
                    if self._parameters.options.compression
                    else zipfile.ZIP_STORED
                )
                buffer = self._repack_xlsx(buffer.getvalue(), compression)
            else:
                workbook = xlwt.Workbook()
                sheet = workbook.add_sheet(sheet_name)
                if len(rows) > 65_536 or any(len(row) > 256 for row in rows):
                    raise UnsupportedNodeConfigurationError(
                        "Формат XLS ограничен 65 536 строками и 256 столбцами."
                    )
                for row_index, row in enumerate(rows):
                    for column_index, value in enumerate(row):
                        if value is None:
                            value = ""
                        if isinstance(value, (date, datetime, time)):
                            value = value.isoformat()
                        if isinstance(value, (dict, list)):
                            value = json.dumps(value, ensure_ascii=False)
                        sheet.write(row_index, column_index, value)
                workbook.save(buffer)
        except UnsupportedNodeConfigurationError:
            raise
        except Exception as error:
            raise UnsupportedNodeConfigurationError(
                "Не удалось создать Excel-файл из входной таблицы."
            ) from error
        if buffer.tell() > _MAX_FILE_BYTES:
            raise UnsupportedNodeConfigurationError("Созданный файл превышает лимит 50 MiB.")
        return buffer.getvalue()

    @staticmethod
    def _repack_xlsx(content: bytes, compression: int) -> io.BytesIO:
        """Set ZIP compression for the generated XLSX package."""
        result = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(content)) as source:
            with zipfile.ZipFile(result, "w", compression=compression) as target:
                for member in source.infolist():
                    target.writestr(member.filename, source.read(member.filename))
        return result


class ExtractFromFileExecutor:
    """Преобразовать CSV, HTML, TXT, Excel, JSON или PDF из binary в JSON items."""

    is_trigger = False

    def __init__(self, parameters: ExtractFileParameters):
        self._parameters = parameters

    async def execute(
        self,
        node: WorkflowNode,
        items: tuple[WorkflowItem, ...],
        context: NodeExecutionContext,
    ) -> NodeResult:
        del node, context
        output: list[WorkflowItem] = []
        paired_input_indexes: list[int] = []
        for item_index, item in enumerate(items):
            content, _ = _binary_bytes(item, self._parameters.binary_property_name)
            operation = self._parameters.operation
            if operation == "csv":
                records = self._extract_csv(content)
            elif operation == "html":
                records = self._extract_html(content)
            elif operation == "fromJson":
                records = self._extract_json(content)
            elif operation == "pdf":
                records = self._extract_pdf(content)
            elif operation == "text":
                records = [self._extract_text(content)]
            elif operation == "binaryToPropery":
                records = [base64.b64encode(content).decode("ascii")]
            elif operation == "xlsx":
                records = self._extract_xlsx(content)
            else:
                records = self._extract_xls(content)

            for record in records:
                keep_source = self._parameters.options.keep_source
                keep_json = keep_source is True or keep_source in {"json", "both"}
                keep_binary = keep_source is True or keep_source in {"binary", "both"}
                preserve_json_source = operation in {"fromJson", "text", "binaryToPropery"}
                json_data = (
                    dict(item.json_data)
                    if preserve_json_source and keep_json and isinstance(item.json_data, Mapping)
                    else {}
                )
                _set_json_path(json_data, self._parameters.destination_key, record)
                binary_data = dict(item.binary or {}) if item.binary else {}
                if not keep_binary:
                    binary_data.pop(self._parameters.binary_property_name, None)
                output.append(
                    WorkflowItem(
                        json=json_data,
                        binary=binary_data or None,
                    )
                )
                paired_input_indexes.append(item_index)
        return NodeResult(
            outputs=(tuple(output),),
            paired_input_indexes=(tuple(paired_input_indexes),),
        )

    def _decode(self, content: bytes) -> str:
        return _decode_file_text(
            content,
            self._parameters.options.encoding,
            self._parameters.options.strip_bom,
        )

    def _extract_csv(self, content: bytes) -> list[dict[str, object]]:
        try:
            reader = csv.reader(
                io.StringIO(self._decode(content)), delimiter=self._parameters.options.delimiter
            )
            rows = []
            for row in reader:
                if len(rows) > _MAX_CSV_ROWS:
                    raise UnsupportedNodeConfigurationError("CSV содержит больше 50 000 строк.")
                rows.append(row)
        except UnsupportedNodeConfigurationError:
            raise
        except csv.Error as error:
            raise UnsupportedNodeConfigurationError("Не удалось прочитать CSV-файл.") from error
        if self._parameters.options.header_row and rows:
            expected_width = len(rows[0])
            if any(len(row) > expected_width for row in rows[1:]):
                raise UnsupportedNodeConfigurationError("В CSV есть строка с лишними столбцами.")
        return _rows_to_records(
            rows,
            header_row=self._parameters.options.header_row,
            include_empty_cells=self._parameters.options.include_empty_cells,
        )

    def _extract_html(self, content: bytes) -> list[dict[str, object]]:
        parser = _FirstHtmlTableParser()
        try:
            parser.feed(self._decode(content))
            parser.close()
        except (ValueError, AssertionError) as error:
            raise UnsupportedNodeConfigurationError("Не удалось прочитать HTML-файл.") from error
        if not parser.rows:
            raise UnsupportedNodeConfigurationError("В HTML-файле не найдена таблица.")
        return _rows_to_records(
            parser.rows,
            header_row=self._parameters.options.header_row,
            include_empty_cells=self._parameters.options.include_empty_cells,
        )

    def _extract_text(self, content: bytes) -> str:
        return self._decode(content)

    def _extract_xlsx(self, content: bytes) -> list[dict[str, object]]:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(member.file_size for member in archive.infolist()) > _MAX_FILE_BYTES * 4:
                    raise UnsupportedNodeConfigurationError(
                        "Распакованный XLSX превышает лимит 200 MiB."
                    )
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            worksheet = (
                workbook[self._parameters.options.sheet_name]
                if self._parameters.options.sheet_name
                else workbook.active
            )
            rows = []
            cell_count = 0
            for row in worksheet.iter_rows(values_only=True):
                cell_count += len(row)
                if len(rows) >= _MAX_CSV_ROWS + 1 or cell_count > _MAX_SPREADSHEET_CELLS:
                    raise UnsupportedNodeConfigurationError(
                        "Таблица XLSX превышает лимит строк или ячеек."
                    )
                rows.append(row)
            workbook.close()
        except UnsupportedNodeConfigurationError:
            raise
        except Exception as error:
            raise UnsupportedNodeConfigurationError("Не удалось прочитать XLSX-файл.") from error
        return _rows_to_records(
            rows,
            header_row=self._parameters.options.header_row,
            include_empty_cells=self._parameters.options.include_empty_cells,
        )

    def _extract_xls(self, content: bytes) -> list[dict[str, object]]:
        try:
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            worksheet = (
                workbook.sheet_by_name(self._parameters.options.sheet_name)
                if self._parameters.options.sheet_name
                else workbook.sheet_by_index(0)
            )
            if worksheet.nrows * worksheet.ncols > _MAX_SPREADSHEET_CELLS:
                raise UnsupportedNodeConfigurationError(
                    "Таблица превышает лимит в 2 000 000 ячеек."
                )
            rows: list[list[object]] = []
            for row_index in range(worksheet.nrows):
                row: list[object] = []
                for cell in worksheet.row(row_index):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode))
                    elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        row.append(None)
                    else:
                        row.append(cell.value)
                rows.append(row)
            workbook.release_resources()
        except UnsupportedNodeConfigurationError:
            raise
        except Exception as error:
            raise UnsupportedNodeConfigurationError("Не удалось прочитать XLS-файл.") from error
        return _rows_to_records(
            rows,
            header_row=self._parameters.options.header_row,
            include_empty_cells=self._parameters.options.include_empty_cells,
        )

    @staticmethod
    def _extract_json(content: bytes) -> list[object]:
        try:
            value = json.loads(content)
        except (UnicodeDecodeError, json.JSONDecodeError) as error:
            raise UnsupportedNodeConfigurationError(
                "Binary-файл не содержит корректный JSON."
            ) from error
        values = value if isinstance(value, list) else [value]
        if any(
            not isinstance(row, (dict, list, str, int, float, bool, type(None))) for row in values
        ):
            raise UnsupportedNodeConfigurationError("JSON содержит неподдерживаемый тип значения.")
        return values

    def _extract_pdf(self, content: bytes) -> list[str]:
        try:
            reader = PdfReader(io.BytesIO(content))
            page_count = len(reader.pages)
            pages = [
                page.extract_text() or ""
                for page in reader.pages[: self._parameters.options.max_pages]
            ]
        except Exception as error:
            raise UnsupportedNodeConfigurationError("Не удалось извлечь текст из PDF.") from error
        if page_count > self._parameters.options.max_pages:
            raise UnsupportedNodeConfigurationError(
                f"PDF превышает лимит {self._parameters.options.max_pages} страниц."
            )
        if self._parameters.options.join_pages:
            return ["\n".join(pages)]
        return pages


def _replace_binary(
    item: WorkflowItem,
    property_name: str,
    binary_value: dict[str, object],
) -> WorkflowItem:
    """Создать item с прежним JSON и добавленным/заменённым binary-полем."""
    binary = dict(item.binary or {})
    binary[property_name] = binary_value
    return WorkflowItem(json=item.json_data, binary=binary)
